from datetime import timedelta
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO, StringIO
from threading import Barrier
from types import SimpleNamespace
from unittest import skipUnless
from unittest.mock import patch

from cryptography.fernet import Fernet
from django.contrib import admin
from django.contrib.messages.storage.fallback import FallbackStorage
from django.core import mail
from django.core.management import call_command
from django.db import connection, connections
from django.test import RequestFactory, override_settings, TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from accounts.admin import UserAdmin
from accounts.models import SkyroomCredential, SkyroomEmail, User
from accounts.skyroom import IMPORT_COLUMNS, credential_password, ensure_credential, export_users, import_room_names, queue_credential_emails
from shop.admin import PresentationAdmin
from shop.models import Participation, Presentation
from shop.serializers import PresentationSerializer


@override_settings(SKYROOM_ENCRYPTION_KEY=Fernet.generate_key().decode(),
                   EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class SkyroomTests(APITestCase):
    url = '/api/users/skyroom/'

    def setUp(self):
        self.user = self.user_record('sky@example.com', '09123456781')
        self.other = self.user_record('other@example.com', '09123456782')
        start = timezone.now() + timedelta(days=10)
        self.presentation = Presentation.objects.create(
            en_title='Online workshop', fa_title='کارگاه آنلاین', cost='10000', capacity=10,
            service_type='WORKSHOP', start=start, end=start + timedelta(hours=1),
            en_description='Test', fa_description='Test', is_skyroom=True,
            skyroom_room_url='https://www.skyroom.online/ch/test/room',
        )
        self.participation = Participation.objects.create(
            user=self.user, presentation=self.presentation, payment_state='COMPLETED',
        )
        self.client.force_authenticate(self.user)

    def user_record(self, email, phone):
        return User.objects.create_user(email=email, phone_number=phone, first_name='Skyroom',
                                        last_name='Tester', password='website-secret-123', is_active=True)

    def test_owner_gets_saved_phone_password_and_room_link(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['username'], 'sky')
        self.assertEqual(response.data['password'], '09123456781')
        self.assertEqual(response.data['presentations'][0]['room_url'], self.presentation.skyroom_room_url)
        self.assertIn('no-store', response['Cache-Control'])
        stored = SkyroomCredential.objects.get(user=self.user)
        self.assertNotIn('09123456781', stored.encrypted_password)
        self.assertEqual(credential_password(stored), response.data['password'])
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password('website-secret-123'))

    def test_credentials_stay_stable_across_panel_export_and_phone_edits(self):
        first = self.client.get(self.url).data
        self.user.phone_number = '09123456789'
        self.user.save()
        again = self.client.get(self.url).data
        response = export_users(User.objects.filter(pk=self.user.pk))
        workbook = load_workbook(BytesIO(response.content))
        row = list(workbook.active.values)[0]
        self.assertEqual(first['password'], again['password'])
        self.assertEqual(row[:2], (first['username'], first['password']))
        self.assertEqual(SkyroomCredential.objects.count(), 1)

    def test_password_removes_only_leading_iran_country_prefix(self):
        self.other.phone_number = '+989123456782'
        self.other.save()
        credential = ensure_credential(self.other)
        self.assertEqual(credential_password(credential), '9123456782')

    def test_login_and_completed_registration_are_required(self):
        self.client.force_authenticate(user=None)
        self.assertEqual(self.client.get(self.url).status_code, 401)
        self.client.force_authenticate(self.other)
        Participation.objects.create(user=self.other, presentation=self.presentation)
        self.assertEqual(self.client.get(self.url).status_code, 404)
        self.assertFalse(SkyroomCredential.objects.exists())

    def test_requested_user_id_cannot_expose_someone_elses_credentials(self):
        self.client.get(self.url)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(self.url, {'user_id': self.user.pk}).status_code, 404)

    def test_room_links_are_not_published_in_public_presentation_response(self):
        data = PresentationSerializer(self.presentation).data
        self.assertTrue(data['is_skyroom'])
        self.assertNotIn('skyroom_room_url', data)

    @override_settings(SKYROOM_ENCRYPTION_KEY='')
    def test_missing_key_returns_service_unavailable_without_creating_credentials(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.data['code'], 'skyroom_not_configured')
        self.assertFalse(SkyroomCredential.objects.exists())

    def test_wrong_key_does_not_replace_existing_password(self):
        self.client.get(self.url)
        encrypted = SkyroomCredential.objects.get().encrypted_password
        with override_settings(SKYROOM_ENCRYPTION_KEY=Fernet.generate_key().decode()):
            self.assertEqual(self.client.get(self.url).status_code, 503)
        self.assertEqual(SkyroomCredential.objects.get().encrypted_password, encrypted)

    def test_username_collisions_are_resolved_without_overwriting_another_user(self):
        first = self.user_record('same@first.example', '09123456783')
        second = self.user_record('same@second.example', '09123456784')
        a, b = ensure_credential(first), ensure_credential(second)
        self.assertEqual(a.username, 'same')
        self.assertEqual(b.username, f'same-{second.pk}')
        self.assertNotEqual(a.username, b.username)
        self.assertNotIn('@', b.username)

    def test_excel_is_headerless_in_skyroom_order_preserves_zeroes_and_blocks_formulas(self):
        self.user.first_name = '=HYPERLINK("evil")'
        self.user.save()
        response = export_users(User.objects.filter(pk=self.user.pk))
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(IMPORT_COLUMNS, ['username', 'password', 'nickname', 'rooms', 'access',
                                          'email', 'mobile', 'firstname', 'lastname'])
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual(sheet.max_column, 9)
        self.assertEqual(sheet['A1'].value, 'sky')
        self.assertEqual(sheet['B1'].value, '09123456781')
        self.assertEqual(sheet['G1'].value, '09123456781')
        self.assertEqual(sheet['C1'].data_type, 's')
        self.assertEqual(sheet['D1'].value, 'room')
        self.assertEqual(sheet['E1'].value, 'normal')
        self.assertEqual(sheet['F1'].value, self.user.email)
        self.assertTrue(all(cell.data_type == 's' for cell in sheet[1]))
        self.assertIn('no-store', response['Cache-Control'])

    def test_export_without_room_omits_access_instead_of_granting_all_rooms(self):
        self.presentation.skyroom_room_url = ''
        self.presentation.save()
        row = next(load_workbook(BytesIO(export_users(User.objects.filter(pk=self.user.pk)).content)).active.values)
        self.assertEqual(row[3:5], (None, None))

    def test_import_rooms_deduplicate_slugs_and_never_grant_wildcard_access(self):
        urls = ['https://www.skyroom.online/ch/account/Room/',
                'https://www.skyroom.online/ch/account/room?x=1',
                'https://www.skyroom.online/ch/account/*',
                'https://www.skyroom.online/ch/account/room%2Cother',
                'https://custom.example/unknown', 'https://[invalid']
        self.assertEqual(import_room_names([SimpleNamespace(skyroom_room_url=url) for url in urls]), 'room')

    def test_export_uses_only_selected_completed_rooms_with_normal_access(self):
        other = Presentation.objects.get(pk=self.presentation.pk)
        other.pk = None
        other.skyroom_room_url = 'https://www.skyroom.online/ch/test/second'
        other.save()
        Participation.objects.create(user=self.user, presentation=other, payment_state='COMPLETED')
        response = export_users(User.objects.filter(pk=self.user.pk), [self.presentation.pk])
        row = next(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(row[3:5], ('room', 'normal'))
        response = export_users(User.objects.filter(pk=self.user.pk))
        row = next(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(row[3:5], ('room,second', 'normal'))

    def test_admin_actions_require_explicit_credential_permissions(self):
        self.other.is_staff = True
        self.other.save()
        request = RequestFactory().get('/admin/shop/presentation/')
        request.user = self.other
        model_admin = PresentationAdmin(Presentation, admin.site)
        self.assertNotIn('export_skyroom_credentials', model_admin.get_actions(request))
        self.assertNotIn('email_skyroom_credentials', model_admin.get_actions(request))
        self.assertNotIn('export_skyroom_credentials', UserAdmin(User, admin.site).get_actions(request))

    def test_presentation_export_only_includes_its_completed_participants(self):
        self.other.is_superuser = True
        self.other.is_staff = True
        self.other.save()
        Participation.objects.create(user=self.other, presentation=self.presentation, payment_state='PENDING')
        request = RequestFactory().post('/admin/shop/presentation/')
        request.user = self.other
        model_admin = PresentationAdmin(Presentation, admin.site)
        response = model_admin.export_skyroom_credentials(request, Presentation.objects.filter(pk=self.presentation.pk))
        rows = list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 'sky')
        self.assertEqual(rows[0][5], self.user.email)

    def test_single_presentation_download_matches_sample_and_excludes_other_orders(self):
        self.other.is_superuser = True
        self.other.is_staff = True
        self.other.save()
        self.presentation.is_skyroom = False
        self.presentation.save()
        self.client.force_authenticate(user=None)
        self.client.force_login(self.other)
        url = reverse('admin:shop_presentation_export_skyroom', args=[self.presentation.pk])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        rows = list(load_workbook(BytesIO(response.content)).active.values)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0], ('sky', '09123456781', 'Skyroom Tester', 'room', 'normal',
                                   self.user.email, self.user.phone_number, 'Skyroom', 'Tester'))

    def test_single_presentation_download_requires_export_permission(self):
        self.other.is_staff = True
        self.other.save()
        self.client.force_authenticate(user=None)
        self.client.force_login(self.other)
        response = self.client.get(reverse('admin:shop_presentation_export_skyroom', args=[self.presentation.pk]))
        self.assertEqual(response.status_code, 403)
        self.assertFalse(SkyroomCredential.objects.exists())

    def test_admin_email_action_requires_confirmation_before_queueing(self):
        self.user.is_staff = True
        self.user.is_superuser = True
        self.user.save()
        factory = RequestFactory()
        request = factory.post('/admin/shop/presentation/', {'action': 'email_skyroom_credentials'})
        request.user = self.user
        model_admin = PresentationAdmin(Presentation, admin.site)
        queryset = Presentation.objects.filter(pk=self.presentation.pk)
        response = model_admin.email_skyroom_credentials(request, queryset)
        self.assertEqual(response.template_name, 'admin/skyroom_email_confirmation.html')
        self.assertEqual(response.context_data['recipient_count'], 1)
        self.assertFalse(SkyroomEmail.objects.exists())
        request = factory.post('/admin/shop/presentation/', {'confirm': 'yes'})
        request.user = self.user
        request.session = {}
        request._messages = FallbackStorage(request)
        model_admin.email_skyroom_credentials(request, queryset)
        self.assertEqual(SkyroomEmail.objects.count(), 1)
        self.assertEqual(len(mail.outbox), 0)

    def test_email_queue_is_private_per_recipient_and_delivery_matches_panel(self):
        count = queue_credential_emails(User.objects.all(), [self.presentation.pk])
        self.assertEqual(count, 1)
        self.assertEqual(len(mail.outbox), 0)
        self.assertNotIn('09123456781', str(SkyroomEmail.objects.get().presentations))
        call_command('send_skyroom_emails', stdout=StringIO())
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, [self.user.email])
        panel = self.client.get(self.url).data
        self.assertIn(panel['username'], mail.outbox[0].body)
        self.assertIn(panel['password'], mail.outbox[0].body)
        self.assertIn(self.presentation.skyroom_room_url, mail.outbox[0].body)
        self.assertEqual(SkyroomEmail.objects.get().status, 'SENT')
        call_command('send_skyroom_emails', stdout=StringIO())
        self.assertEqual(len(mail.outbox), 1)

    @patch('accounts.management.commands.send_skyroom_emails.send_mail', side_effect=RuntimeError('SMTP unavailable'))
    def test_email_failure_retries_without_resetting_credentials(self, send):
        queue_credential_emails(User.objects.filter(pk=self.user.pk), [self.presentation.pk])
        original = SkyroomCredential.objects.get().encrypted_password
        for index in range(5):
            SkyroomEmail.objects.update(available_at=timezone.now() - timedelta(seconds=1))
            call_command('send_skyroom_emails', stdout=StringIO())
        job = SkyroomEmail.objects.get()
        self.assertEqual(job.status, 'FAILED')
        self.assertEqual(job.attempts, 5)
        self.assertEqual(job.last_error, 'RuntimeError')
        self.assertEqual(SkyroomCredential.objects.get().encrypted_password, original)

    def test_revoked_registration_is_not_emailed(self):
        queue_credential_emails(User.objects.filter(pk=self.user.pk), [self.presentation.pk])
        self.participation.delete()
        call_command('send_skyroom_emails', stdout=StringIO())
        self.assertEqual(len(mail.outbox), 0)
        self.assertEqual(SkyroomEmail.objects.get().status, 'FAILED')

    def test_email_address_change_does_not_send_credentials_to_old_recipient(self):
        queue_credential_emails(User.objects.filter(pk=self.user.pk), [self.presentation.pk])
        self.user.email = 'updated@example.com'
        self.user.save()
        call_command('send_skyroom_emails', stdout=StringIO())
        self.assertEqual(len(mail.outbox), 0)
        self.assertEqual(SkyroomEmail.objects.get().last_error, 'RecipientChangedOrInactive')


@skipUnless(connection.vendor == 'postgresql', 'Credential concurrency requires PostgreSQL.')
@override_settings(SKYROOM_ENCRYPTION_KEY=Fernet.generate_key().decode())
class SkyroomConcurrencyTests(TransactionTestCase):
    def test_simultaneous_panel_and_export_creation_reuses_one_saved_login(self):
        user = User.objects.create_user(email='concurrent@example.com', phone_number='09123456780',
                                        first_name='Concurrent', last_name='User', password='website-secret')
        barrier = Barrier(2)

        def execute(_):
            try:
                barrier.wait(timeout=10)
                record = ensure_credential(user)
                return record.pk, record.username, credential_password(record)
            finally:
                connections.close_all()

        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(execute, range(2)))
        self.assertEqual(results[0], results[1])
        self.assertEqual(SkyroomCredential.objects.count(), 1)
